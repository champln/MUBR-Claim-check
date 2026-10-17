"""
Pre-screen สด (กึ่ง real-time) — ดึง visit จากฐาน HOSxP โดยตรงแล้วรันกฎตรวจสอบทันที
ไม่บันทึกผลลง DB (ephemeral) — frontend จะ poll ซ้ำตามรอบเวลาที่ผู้ใช้ตั้ง
"""
from __future__ import annotations

import io
import threading
import time
from datetime import date, datetime, timedelta
from typing import Any, Optional

from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from auth import get_current_user
from database import get_db
from models import HosxpConnectionConfig, User, ValidationRule
from services.claim_checker import run_prescreen
from services.hosxp_live_source import (
    HosxpLiveSourceError,
    check_readiness,
    fetch_live_opd_records,
    list_pttypes,
    test_connection,
)

router = APIRouter(prefix="/live-prescreen", tags=["live-prescreen"])

# แคชผลรอบล่าสุดต่อเงื่อนไข (ช่วงวันที่ + สิทธิ) — หลายเครื่องเปิดหน้าพร้อมกัน
# จะไม่ยิง query เข้าฐาน HOSxP ซ้ำๆ ภายในช่วงนี้ (ปุ่ม "ดึงข้อมูลตอนนี้" ข้ามแคชได้)
CACHE_TTL_SECONDS = 30
_cache: dict[tuple, tuple[float, dict]] = {}
_cache_lock = threading.Lock()


def _get_db_url(db: Session) -> Optional[str]:
    cfg = db.query(HosxpConnectionConfig).filter(HosxpConnectionConfig.is_enabled == True).first()  # noqa: E712
    return cfg.db_url if cfg else None


@router.get("/status")
def live_status(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """เช็คว่าตั้งค่าเชื่อม HOSxP แล้วหรือยัง และเชื่อมได้จริงไหม"""
    db_url = _get_db_url(db)
    if not db_url:
        return {"configured": False, "reachable": False,
                "message": "ยังไม่ได้ตั้งค่าเชื่อมต่อฐาน HOSxP (ADMIN ตั้งได้ที่ ตั้งค่าระบบ → การเชื่อมต่อฐาน HOSxP)"}
    try:
        test_connection(db_url)
        return {"configured": True, "reachable": True, "message": "เชื่อมต่อ HOSxP ได้"}
    except HosxpLiveSourceError as e:
        return {"configured": True, "reachable": False, "message": str(e)}


@router.get("/readiness")
def live_readiness(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """ตรวจว่าฐาน HOSxP มีตาราง/คอลัมน์ครบสำหรับกฎแต่ละข้อ (อ่านเฉพาะโครงสร้าง)"""
    db_url = _get_db_url(db)
    if not db_url:
        raise HTTPException(400, "ยังไม่ได้ตั้งค่าเชื่อมต่อฐาน HOSxP")
    try:
        return check_readiness(db_url)
    except HosxpLiveSourceError as e:
        raise HTTPException(502, str(e))


@router.get("")
def live_prescreen(
    date_from: Optional[date] = Query(None, description="วันที่เริ่ม (default: วันนี้)"),
    date_to: Optional[date] = Query(None, description="วันที่สิ้นสุด (default: วันนี้)"),
    pttype: Optional[str] = Query(None, description="รหัสสิทธิ HOSxP คั่นด้วย comma เช่น 21,10"),
    limit: int = Query(500, ge=1, le=2000),
    force: bool = Query(False, description="ข้ามแคช ดึงจาก HOSxP ใหม่ทันที"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """ดึง visit จาก HOSxP แล้วรัน pre-screen ทันที — คืนผลโดยไม่บันทึก"""
    db_url = _get_db_url(db)
    if not db_url:
        raise HTTPException(400, "ยังไม่ได้ตั้งค่าเชื่อมต่อฐาน HOSxP")

    today = date.today()
    d1 = date_from or today
    d2 = date_to or today
    if d2 < d1:
        raise HTTPException(400, "ช่วงวันที่ไม่ถูกต้อง")
    if (d2 - d1) > timedelta(days=31):
        raise HTTPException(400, "ช่วงวันที่ต้องไม่เกิน 31 วัน (โหมดสดออกแบบไว้ตรวจข้อมูลล่าสุด)")

    pttypes = [p.strip() for p in pttype.split(",") if p.strip()] if pttype else None

    key = (str(d1), str(d2), tuple(sorted(pttypes or [])), limit)
    if not force:
        with _cache_lock:
            hit = _cache.get(key)
        if hit and time.monotonic() - hit[0] < CACHE_TTL_SECONDS:
            return {**hit[1], "cached": True}

    try:
        records = fetch_live_opd_records(db_url, d1, d2, pttypes, limit)
    except HosxpLiveSourceError as e:
        raise HTTPException(502, str(e))

    # กฎ custom ที่เปิดใช้อยู่ — ชุดเดียวกับ pre-screen ปกติ
    custom_rules = [
        {
            "rule_code": r.rule_code,
            "rule_name": r.rule_name,
            "rule_name_th": r.rule_name_th,
            "category": r.category,
            "condition_type": r.condition_type,
            "condition_value": r.condition_value,
            "severity": r.severity.value,
            "is_active": r.is_active,
        }
        for r in db.query(ValidationRule).filter(ValidationRule.is_active == True).all()  # noqa: E712
    ]

    results: list[dict[str, Any]] = []
    passed = failed = flagged_c = warned = 0
    total_amount = 0.0

    for rec in records:
        errors, is_c, flag_reason = run_prescreen(rec, custom_rules)
        has_error = any(e.severity == "ERROR" for e in errors)
        has_warning = any(e.severity == "WARNING" for e in errors)

        if is_c:
            status = "FLAGGED_C"
            flagged_c += 1
        elif has_error:
            status = "FAILED"
            failed += 1
        else:
            status = "PASSED"
            passed += 1
        if has_warning:
            warned += 1
        total_amount += rec.get("total_charge") or 0.0

        results.append({
            "vn": rec.get("claim_no"),
            "hn": rec.get("hn"),
            "pid": rec.get("pid"),
            "patient_name": rec.get("patient_name"),
            "visit_date": rec.get("visit_date"),
            "visit_time": rec.get("vsttime"),
            "pttype": rec.get("pttype"),
            "spclty_name": rec.get("spclty_name") or rec.get("spclty") or "",
            "main_dep": rec.get("main_dep") or "",
            "main_dep_name": rec.get("main_dep_name") or rec.get("main_dep") or "",
            "cur_dep_name": rec.get("cur_dep_name") or rec.get("cur_dep") or "",
            "pttype_name": rec.get("pttype_name") or "",
            "doctor_name": rec.get("doctor_name") or "",
            "pdx": rec.get("pdx"),
            "adx": [rec.get(f"adx{i}") for i in range(1, 5) if rec.get(f"adx{i}")],
            "total_charge": rec.get("total_charge") or 0.0,
            "status": status,
            "flag_reason": flag_reason if is_c else None,
            "errors": [
                {
                    "code": e.error_code,
                    "category": e.error_category,
                    "field": e.error_field,
                    "message_th": e.error_message_th,
                    "severity": e.severity,
                    "current_value": e.current_value,
                }
                for e in errors
            ],
        })

    result = {
        "fetched_at": datetime.now().isoformat(timespec="seconds"),
        "date_from": str(d1),
        "date_to": str(d2),
        "pttype": pttypes,
        "summary": {
            "total": len(results),
            "passed": passed,
            "failed": failed,
            "flagged_c": flagged_c,
            "warned": warned,
            "total_amount": round(total_amount, 2),
        },
        "records": results,
        # กฎจากเมนูแก้ไฟล์ที่รันกับ HOSxP นี้ไม่ได้ (เช่นชื่อคอลัมน์ต่างจากมาตรฐาน)
        "rules_skipped": records[0].get("file_rule_skipped", []) if records else [],
        "cache_ttl": CACHE_TTL_SECONDS,
    }
    with _cache_lock:
        # เก็บเฉพาะรอบล่าสุดของแต่ละเงื่อนไข และไม่ให้แคชโตไม่จำกัด
        if len(_cache) > 50:
            _cache.clear()
        _cache[key] = (time.monotonic(), result)
    return {**result, "cached": False}


@router.get("/pttypes")
def live_pttypes(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """รายชื่อสิทธิการรักษาจาก HOSxP (ทะเบียน pttype) สำหรับ dropdown ตัวกรองสิทธิ"""
    db_url = _get_db_url(db)
    if not db_url:
        raise HTTPException(400, "ยังไม่ได้ตั้งค่าเชื่อมต่อฐาน HOSxP")
    try:
        return list_pttypes(db_url)
    except HosxpLiveSourceError as e:
        raise HTTPException(502, str(e))


SEVERITY_TH = {"ERROR": "ข้อผิดพลาด", "WARNING": "เตือน", "INFO": "แจ้งให้ทราบ"}
STATUS_TH = {"PASSED": "ผ่าน", "FAILED": "ไม่ผ่าน", "FLAGGED_C": "เสี่ยงติด C"}


@router.post("/export")
def live_export(
    payload: dict = Body(..., description="{ records: [...ผล pre-screen สด], title?: str }"),
    current_user: User = Depends(get_current_user),
):
    """
    ส่งออกผล Pre-screen สด (ตามตัวกรองที่เลือกอยู่บนหน้าเว็บ) เป็น Excel
    ชีต 1 รายการปัญหา (1 แถวต่อ 1 ปัญหา) · ชีต 2 สรุปตามรหัส · ชีต 3 สรุปตามแผนก · ชีต 4 visit ทั้งหมด
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    records = payload.get("records") or []
    if not isinstance(records, list) or len(records) > 20000:
        raise HTTPException(400, "ข้อมูลไม่ถูกต้อง")
    title = str(payload.get("title") or "Pre-screen สด ผู้ป่วยนอก")[:120]

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E78")
    sev_fill = {"ERROR": "F8CBAD", "WARNING": "FFE699", "INFO": "DDEBF7"}

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.data_type == "f":   # ข้อความที่ขึ้นต้นด้วย = ให้เป็นข้อความ ไม่ใช่สูตร
                    c.data_type = "s"
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    def s(v: Any) -> Optional[str]:
        return None if v is None or v == "" else str(v)

    def base(r: dict) -> list:
        return [s(r.get("vn")), s(r.get("hn")), s(r.get("patient_name")),
                s(r.get("visit_date")), s(str(r.get("visit_time") or "")[:5]),
                s(" ".join(x for x in (r.get("pttype"), r.get("pttype_name")) if x)),
                s(r.get("spclty_name")), s(r.get("main_dep_name")), s(r.get("doctor_name"))]

    base_head = ["VN", "HN", "ชื่อผู้ป่วย", "วันที่", "เวลา", "สิทธิ", "แผนก", "ห้องตรวจ", "แพทย์ผู้ตรวจ"]
    base_w = [14, 10, 24, 11, 7, 26, 18, 20, 22]

    # 1) รายการปัญหา
    ws = wb.active
    ws.title = "รายการปัญหา"
    issue_rows = []
    by_code: dict[str, dict] = {}
    by_dept: dict[str, dict] = {}
    for r in records:
        dept = r.get("spclty_name") or "(ไม่ระบุ)"
        d = by_dept.setdefault(dept, {"total": 0, "bad": 0, "err": 0, "warn": 0, "codes": {}})
        d["total"] += 1
        if r.get("status") != "PASSED":
            d["bad"] += 1
        for e in r.get("errors") or []:
            code = str(e.get("code") or "")
            sev = str(e.get("severity") or "")
            issue_rows.append(base(r) + [s(code), SEVERITY_TH.get(sev, sev), s(e.get("message_th")),
                                         s(e.get("current_value")), STATUS_TH.get(r.get("status"), r.get("status"))])
            c = by_code.setdefault(code, {"sev": sev, "n": 0, "vns": set(), "example": e.get("message_th") or ""})
            c["n"] += 1
            c["vns"].add(r.get("vn"))
            if sev in ("ERROR", "WARNING"):
                d["err" if sev == "ERROR" else "warn"] += 1
            d["codes"][code] = d["codes"].get(code, 0) + 1
    sheet(ws, base_head + ["รหัส", "ระดับ", "รายละเอียด", "ค่าปัจจุบัน", "สถานะ visit"], issue_rows,
          base_w + [9, 11, 70, 18, 12])
    for row in ws.iter_rows(min_row=2):
        sev = next((k for k, v in SEVERITY_TH.items() if v == row[10].value), None)
        if sev:
            row[10].fill = PatternFill("solid", fgColor=sev_fill[sev])

    # 2) สรุปตามรหัส
    code_rows = sorted(by_code.items(), key=lambda kv: (-len(kv[1]["vns"]), kv[0]))
    sheet(wb.create_sheet("สรุปตามรหัส"), ["รหัส", "ระดับ", "จำนวน visit", "จำนวนปัญหา", "ตัวอย่างข้อความ"],
          [[s(k), SEVERITY_TH.get(v["sev"], v["sev"]), len(v["vns"]), v["n"], s(v["example"])] for k, v in code_rows],
          [10, 12, 12, 12, 90])

    # 3) สรุปตามแผนก
    dept_rows = sorted(by_dept.items(), key=lambda kv: (-kv[1]["bad"], -kv[1]["total"]))
    sheet(wb.create_sheet("สรุปตามแผนก"),
          ["แผนก", "visit ทั้งหมด", "visit ที่ไม่ผ่าน/เสี่ยง C", "ข้อผิดพลาด", "เตือน", "รหัสที่พบบ่อย"],
          [[s(k), v["total"], v["bad"], v["err"], v["warn"],
            ", ".join(f"{c} ({n})" for c, n in sorted(v["codes"].items(), key=lambda x: -x[1])[:6])]
           for k, v in dept_rows],
          [26, 13, 18, 12, 10, 60])

    # 4) visit ทั้งหมด
    sheet(wb.create_sheet("visit ทั้งหมด"),
          base_head + ["PDX", "ค่าใช้จ่าย (บาท)", "สถานะ", "ข้อผิดพลาด", "เตือน", "รหัสปัญหา"],
          [base(r) + [s(r.get("pdx")), float(r.get("total_charge") or 0), STATUS_TH.get(r.get("status"), r.get("status")),
                      sum(1 for e in r.get("errors") or [] if e.get("severity") == "ERROR"),
                      sum(1 for e in r.get("errors") or [] if e.get("severity") == "WARNING"),
                      ", ".join(sorted({str(e.get("code")) for e in r.get("errors") or []}))]
           for r in records],
          base_w + [9, 14, 12, 11, 8, 30])

    for w in wb.worksheets:
        w.sheet_properties.tabColor = "1F4E78"
    wb.properties.title = title

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"live_prescreen_opd_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ─── ผู้ป่วยใน ─────────────────────────────────────────────────────────────────
# ดึงเฉพาะเมื่อเปิดหน้าผู้ป่วยใน (หน้า OPD ไม่เรียก endpoint นี้)

@router.get("/ipd")
def live_prescreen_ipd(
    mode: str = Query("discharged", pattern="^(discharged|admitted)$",
                      description="discharged = จำหน่ายแล้วในช่วงวันที่ · admitted = ยังนอนอยู่"),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    limit: int = Query(300, ge=1, le=1000),
    force: bool = Query(False),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from services.hosxp_live_ipd import DISCHARGE_ONLY_CODES, fetch_live_ipd_records

    db_url = _get_db_url(db)
    if not db_url:
        raise HTTPException(400, "ยังไม่ได้ตั้งค่าเชื่อมต่อฐาน HOSxP")
    today = date.today()
    d1 = date_from or today
    d2 = date_to or today
    if mode == "discharged":
        if d2 < d1:
            raise HTTPException(400, "ช่วงวันที่ไม่ถูกต้อง")
        if (d2 - d1) > timedelta(days=31):
            raise HTTPException(400, "ช่วงวันที่ต้องไม่เกิน 31 วัน")

    key = ("ipd", mode, str(d1) if mode == "discharged" else "", str(d2) if mode == "discharged" else "", limit)
    if not force:
        with _cache_lock:
            hit = _cache.get(key)
        if hit and time.monotonic() - hit[0] < CACHE_TTL_SECONDS:
            return {**hit[1], "cached": True}

    try:
        records = fetch_live_ipd_records(db_url, mode, d1, d2, limit)
    except HosxpLiveSourceError as e:
        raise HTTPException(502, str(e))

    custom_rules = [
        {
            "rule_code": r.rule_code, "rule_name": r.rule_name, "rule_name_th": r.rule_name_th,
            "category": r.category, "condition_type": r.condition_type,
            "condition_value": r.condition_value, "severity": r.severity.value, "is_active": r.is_active,
        }
        for r in db.query(ValidationRule).filter(ValidationRule.is_active == True).all()  # noqa: E712
    ]

    results: list[dict[str, Any]] = []
    passed = failed = flagged_c = warned = 0
    total_amount = 0.0
    for rec in records:
        errors, is_c, flag_reason = run_prescreen(rec, custom_rules)
        if mode == "admitted":
            # ยังนอนอยู่ — ยังไม่มีวันจำหน่าย/DRG/PDX เป็นเรื่องปกติ
            errors = [e for e in errors if e.error_code not in DISCHARGE_ONLY_CODES]
        has_error = any(e.severity == "ERROR" for e in errors)
        if is_c:
            status, flagged_c = "FLAGGED_C", flagged_c + 1
        elif has_error:
            status, failed = "FAILED", failed + 1
        else:
            status, passed = "PASSED", passed + 1
        if any(e.severity == "WARNING" for e in errors):
            warned += 1
        total_amount += rec.get("total_charge") or 0.0
        results.append({
            "an": rec["an"],
            "vn": rec.get("vn"),
            "hn": rec.get("hn"),
            "pid": rec.get("pid"),
            "patient_name": rec.get("patient_name"),
            "admit_date": rec.get("visit_date"),
            "admit_time": rec.get("regtime"),
            "discharge_date": rec.get("discharge_date"),
            "discharge_time": rec.get("dchtime"),
            "los": rec.get("los"),
            "pttype": rec.get("pttype"),
            "pttype_name": rec.get("pttype_name") or "",
            "ward_name": rec.get("ward_name") or "",
            "spclty_name": rec.get("spclty_name") or "",
            "admit_doctor_name": rec.get("admit_doctor_name") or "",
            "dch_doctor_name": rec.get("dch_doctor_name") or "",
            "pdx": rec.get("pdx"),
            "adx": [rec.get(f"adx{i}") for i in range(1, 5) if rec.get(f"adx{i}")],
            "drg": rec.get("drg_code"),
            "rw": rec.get("rw"),
            "auth_code": rec.get("auth_code") or "",
            "total_charge": rec.get("total_charge") or 0.0,
            "status": status,
            "flag_reason": flag_reason if is_c else None,
            "errors": [
                {"code": e.error_code, "category": e.error_category, "field": e.error_field,
                 "message_th": e.error_message_th, "severity": e.severity, "current_value": e.current_value}
                for e in errors
            ],
        })

    result = {
        "fetched_at": datetime.now().isoformat(timespec="seconds"),
        "mode": mode,
        "date_from": str(d1),
        "date_to": str(d2),
        "summary": {"total": len(results), "passed": passed, "failed": failed,
                    "flagged_c": flagged_c, "warned": warned, "total_amount": round(total_amount, 2)},
        "records": results,
        "cache_ttl": CACHE_TTL_SECONDS,
    }
    with _cache_lock:
        if len(_cache) > 50:
            _cache.clear()
        _cache[key] = (time.monotonic(), result)
    return {**result, "cached": False}


@router.get("/claimcat-master")
def claimcat_master(
    days: int = Query(180, ge=7, le=730, description="นับการใช้กับผู้ป่วยในย้อนหลังกี่วัน"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """รายงานทะเบียนรายการ HOSxP ที่ ClaimCat น่าสงสัย (ส่งให้ BMS/IT แก้ที่ต้นทาง)"""
    from services.hosxp_live_ipd import claimcat_master_report

    db_url = _get_db_url(db)
    if not db_url:
        raise HTTPException(400, "ยังไม่ได้ตั้งค่าเชื่อมต่อฐาน HOSxP")
    try:
        return claimcat_master_report(db_url, date.today() - timedelta(days=days))
    except HosxpLiveSourceError as e:
        raise HTTPException(502, str(e))
